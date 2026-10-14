from pathlib import Path
from openpyxl import load_workbook

file_path = Path(
    "data/raw/01_LedgerForge_Python_Master_Dataset.xlsx"
)

workbook = load_workbook(
    file_path,
    read_only=True,
    data_only=True
)

sheet = workbook["employees"]

print("EMPLOYEES SHEET")
print("=" * 60)

for row in sheet.iter_rows(min_row=1, max_row=11, values_only=True):
    print(row)

workbook.close()